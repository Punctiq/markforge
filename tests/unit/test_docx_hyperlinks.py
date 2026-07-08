from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from app.converters.docx_converter import DocxConverter
from app.quality.markdown_quality import build_quality_report


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    if text:
        run = OxmlElement("w:r")
        text_el = OxmlElement("w:t")
        text_el.text = text
        run.append(text_el)
        hyperlink.append(run)
    paragraph._p.append(hyperlink)




def _add_anchor_hyperlink(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _convert_docx(path: Path) -> str:
    return DocxConverter(config={}).convert(path).markdown


def test_docx_table_cell_hyperlink_preserves_text_and_target(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Document type"
    table.cell(0, 1).text = "Document name"
    table.cell(1, 0).text = "Deployment Guide"
    _add_hyperlink(table.cell(1, 1).paragraphs[0], "Project Wiki", "https://dev.example.test/wiki/guide")
    path = tmp_path / "table-link.docx"
    doc.save(path)

    markdown = _convert_docx(path)

    assert "| Deployment Guide | [Project Wiki](https://dev.example.test/wiki/guide) |" in markdown


def test_docx_paragraph_hyperlink_preserves_text_and_target(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("See ")
    _add_hyperlink(paragraph, "Runbook", "https://example.test/runbook")
    paragraph.add_run(" for details.")
    path = tmp_path / "paragraph-link.docx"
    doc.save(path)

    markdown = _convert_docx(path)

    assert "See [Runbook](https://example.test/runbook) for details." in markdown


def test_docx_hyperlink_empty_display_text_falls_back_to_url(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Reference: ")
    _add_hyperlink(paragraph, "", "https://example.test/empty-display")
    path = tmp_path / "empty-display-link.docx"
    doc.save(path)

    markdown = _convert_docx(path)

    assert "Reference: [https://example.test/empty-display](https://example.test/empty-display)" in markdown


def test_quality_report_flags_poor_hyperlink_target_retention():
    original = "[A](https://example.test/a) [B](https://example.test/b) [C](https://example.test/c)"
    final = "A B [C](https://example.test/c)"

    report = build_quality_report(
        original_markdown=original,
        final_markdown=final,
        mode="safe",
        ai_cleanup_requested=False,
        ai_applied=False,
        config=None,
    )

    findings = report["deterministic_report"]["findings"]
    assert report["integrity_risk_level"] == "medium"
    assert report["recommendation"] == "review_manually"
    assert any(item["type"] == "missing_hyperlinks_or_link_targets" for item in findings)
    assert report["deterministic_report"]["metrics"]["hyperlink_target_retention_ratio"] == 1 / 3

def test_docx_normal_non_hyperlink_paragraphs_are_unchanged(tmp_path):
    doc = Document()
    doc.add_paragraph("First plain paragraph.")
    doc.add_paragraph("Second plain paragraph with punctuation: alpha, beta.")
    path = tmp_path / "plain-paragraphs.docx"
    doc.save(path)

    markdown = _convert_docx(path)

    assert markdown == "First plain paragraph.\nSecond plain paragraph with punctuation: alpha, beta."


def test_docx_hyperlink_traversal_preserves_run_order(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("Alpha ")
    _add_hyperlink(paragraph, "one", "https://example.test/one")
    paragraph.add_run(" beta ")
    _add_hyperlink(paragraph, "two", "https://example.test/two")
    paragraph.add_run(" omega")
    path = tmp_path / "ordered-links.docx"
    doc.save(path)

    markdown = _convert_docx(path)

    assert markdown == "Alpha [one](https://example.test/one) beta [two](https://example.test/two) omega"


def test_docx_hyperlink_targets_cover_common_schemes_and_internal_anchors(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph()
    _add_hyperlink(paragraph, "HTTP", "http://example.test/path")
    paragraph.add_run(" ")
    _add_hyperlink(paragraph, "HTTPS", "https://example.test/path")
    paragraph.add_run(" ")
    _add_hyperlink(paragraph, "Mail", "mailto:team@example.test")
    paragraph.add_run(" ")
    _add_hyperlink(paragraph, "Relative", "../docs/page.md")
    paragraph.add_run(" ")
    _add_anchor_hyperlink(paragraph, "Section", "section_1")
    path = tmp_path / "target-types.docx"
    doc.save(path)

    markdown = _convert_docx(path)

    assert markdown == (
        "[HTTP](http://example.test/path) "
        "[HTTPS](https://example.test/path) "
        "[Mail](mailto:team@example.test) "
        "[Relative](../docs/page.md) "
        "[Section](#section_1)"
    )


def test_docx_hyperlink_escapes_link_text_and_unsafe_target_characters(tmp_path):
    doc = Document()
    paragraph = doc.add_paragraph("See ")
    _add_hyperlink(
        paragraph,
        r"Use [draft] \ path",
        "https://example.test/path (draft)/file?q=a(b)",
    )
    path = tmp_path / "escaped-link.docx"
    doc.save(path)

    markdown = _convert_docx(path)

    assert markdown == (
        r"See [Use \[draft\] \\ path]"
        "(https://example.test/path%20%28draft%29/file?q=a%28b%29)"
    )


def test_quality_report_has_neutral_hyperlink_retention_when_source_has_no_links():
    report = build_quality_report(
        original_markdown="Plain source text.",
        final_markdown="Plain source text.",
        mode="safe",
        ai_cleanup_requested=False,
        ai_applied=False,
        config=None,
    )

    findings = report["deterministic_report"]["findings"]
    assert report["integrity_risk_level"] == "low"
    assert report["deterministic_report"]["metrics"]["hyperlink_target_retention_ratio"] == 1.0
    assert not any(item["type"] == "missing_hyperlinks_or_link_targets" for item in findings)


def test_quality_report_keeps_low_integrity_risk_when_hyperlink_targets_are_retained():
    report = build_quality_report(
        original_markdown="[Alpha](https://example.test/a) [Beta](mailto:team@example.test)",
        final_markdown="[Renamed](https://example.test/a) [Contact](mailto:team@example.test)",
        mode="safe",
        ai_cleanup_requested=False,
        ai_applied=False,
        config=None,
    )

    findings = report["deterministic_report"]["findings"]
    assert report["integrity_risk_level"] == "low"
    assert report["deterministic_report"]["metrics"]["hyperlink_target_retention_ratio"] == 1.0
    assert not any(item["type"] == "missing_hyperlinks_or_link_targets" for item in findings)

